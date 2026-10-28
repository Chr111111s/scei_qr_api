from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.entities import ReportData
from app.domain.ports.exports import XlsxExporterPort


class XlsxExporter(XlsxExporterPort):
    def build(self, data: ReportData) -> bytes:
        workbook = Workbook()
        workbook.creator = "scei_qr_api"

        equipment_sheet = workbook.active
        equipment_sheet.title = "Equipos"
        self._write_equipment(equipment_sheet, data)
        self._style_header(equipment_sheet)

        usage_sheet = workbook.create_sheet("Usos")
        self._write_usages(usage_sheet, data)
        self._style_header(usage_sheet)

        incident_sheet = workbook.create_sheet("Incidencias")
        self._write_incidents(incident_sheet, data)
        self._style_header(incident_sheet)

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _write_equipment(sheet: Worksheet, data: ReportData) -> None:
        sheet.append(["Código", "Nombre", "Ubicación", "Activo", "Usos", "Incidencias"])
        for item in data.equipos:
            equipo = item.equipment
            sheet.append(
                [
                    equipo.codigo,
                    equipo.nombre,
                    equipo.ubicacion,
                    "Sí" if equipo.activo else "No",
                    item.usos,
                    item.incidencias,
                ]
            )
        for column, width in zip("ABCDEF", [15, 30, 30, 10, 10, 12], strict=True):
            sheet.column_dimensions[column].width = width

    @staticmethod
    def _write_usages(sheet: Worksheet, data: ReportData) -> None:
        sheet.append(["Alumno", "Equipo", "Inicio", "Fin", "Estado"])
        for usage in data.usos:
            alumno = usage.alumno.full_name if usage.alumno else "-"
            equipo = usage.equipo.codigo if usage.equipo else "-"
            sheet.append(
                [
                    alumno,
                    equipo,
                    usage.inicio,
                    usage.fin or "-",
                    usage.estado.value,
                ]
            )
        for column, width in zip("ABCDE", [30, 20, 22, 22, 12], strict=True):
            sheet.column_dimensions[column].width = width

    @staticmethod
    def _write_incidents(sheet: Worksheet, data: ReportData) -> None:
        sheet.append(
            [
                "Equipo",
                "Título",
                "Reportada por",
                "Atendida por",
                "Estado",
                "Solución",
                "Creada",
            ]
        )
        for incident in data.incidencias:
            sheet.append(
                [
                    incident.equipo.codigo if incident.equipo else "-",
                    incident.titulo,
                    incident.reportada_por.full_name if incident.reportada_por else "-",
                    incident.atendida_por.full_name if incident.atendida_por else "-",
                    incident.estado.value,
                    incident.solucion or "-",
                    incident.creada_en,
                ]
            )
        for column, width in zip("ABCDEFG", [15, 35, 25, 25, 15, 40, 22], strict=True):
            sheet.column_dimensions[column].width = width

    @staticmethod
    def _style_header(sheet: Worksheet) -> None:
        header_fill = PatternFill(fill_type="solid", fgColor="FF2F54EB")
        header_font = Font(bold=True, color="FFFFFFFF")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        sheet.freeze_panes = "A2"
